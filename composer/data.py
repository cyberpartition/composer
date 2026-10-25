"""Load spreadsheets and fill bracket tokens such as [First Name]."""

from __future__ import annotations

import csv
import io
import math
import re
from dataclasses import dataclass
from datetime import date, datetime, time
from pathlib import Path
from zipfile import BadZipFile

EMAIL = "email"
WORD = "word"

EXAMPLE_SUBJECT = "Hello [First Name]"
EXAMPLE_BODY = (
    "Hello [First Name] [Last Name],\n"
    "\n"
    "We noticed your phone number is [Phone] and your email is [Email]. Is this correct?"
)

_TOKEN = re.compile(r"\[([^\[\]\r\n]+)\]")
_INVALID_FILENAME = re.compile(r'[<>:"/\\|?*\x00-\x1f\x7f]')
_RESERVED = {"CON", "PRN", "AUX", "NUL", *(f"COM{i}" for i in range(1, 10)), *(f"LPT{i}" for i in range(1, 10))}
_EMAIL_NAMES = ("email", "e-mail", "email address", "e-mail address", "mail", "emailaddress")


@dataclass
class Rendered:
    text: str
    unknown: list[str]


@dataclass
class Table:
    headers: list[str]
    records: list[dict[str, str]]
    sheets: list[str]
    sheet: str | None
    source: Path


def normalize(value: str) -> str:
    return re.sub(r"\s+", " ", value).strip().casefold()


def single_line(value: str) -> str:
    return re.sub(r"[\r\n]+", " ", value).strip()


def find_header(headers: list[str], names: tuple[str, ...]) -> str | None:
    lookup = {normalize(header): header for header in headers}
    for name in names:
        found = lookup.get(normalize(name))
        if found:
            return found
    return None


def guess_email_column(headers: list[str]) -> str | None:
    found = find_header(headers, _EMAIL_NAMES)
    if found:
        return found
    for header in headers:
        folded = normalize(header)
        if "email" in folded or "e-mail" in folded:
            return header
    return None


def closest_header(name: str, headers: list[str]) -> str | None:
    key = normalize(name)
    if len(key) < 3:
        return None
    for header in headers:
        folded = normalize(header)
        if len(folded) >= 3 and (key in folded or folded in key):
            return header
    return None


def looks_like_email(value: str) -> bool:
    value = value.strip()
    if not value or len(value) > 254 or any(character.isspace() for character in value):
        return False
    if value.count("@") != 1:
        return False
    local, domain = value.split("@")
    if not local or not domain or "." not in domain:
        return False
    if domain.startswith(".") or domain.endswith(".") or ".." in domain:
        return False
    return True


def render(
    template: str,
    record: dict[str, str],
    headers: list[str],
    row_number: int,
    extras: dict[str, str] | None = None,
) -> Rendered:
    lookup = {normalize(header): header for header in headers}
    extra_lookup = {normalize(key): value for key, value in (extras or {}).items()}
    unknown: list[str] = []
    seen: set[str] = set()

    def replace(match: re.Match[str]) -> str:
        token = match.group(1)
        key = normalize(token)
        if key in lookup:
            return record.get(lookup[key], "")
        if key in extra_lookup:
            return extra_lookup[key]
        if key == "row" and "row" not in lookup:
            return str(row_number)
        raw = match.group(0)
        if raw not in seen:
            seen.add(raw)
            unknown.append(raw)
        return raw

    return Rendered(_TOKEN.sub(replace, template), unknown)


_ZIP_NAMES = ("zip", "zip code", "zipcode", "postal code", "postalcode", "postal", "zip_code", "zcta")


def guess_zip_column(headers: list[str]) -> str | None:
    found = find_header(headers, _ZIP_NAMES)
    if found:
        return found
    for header in headers:
        folded = normalize(header)
        if "zip" in folded or folded.startswith("postal"):
            return header
    return None


def default_filename_pattern(headers: list[str]) -> str:
    first = find_header(headers, ("first name", "firstname", "given name"))
    last = find_header(headers, ("last name", "lastname", "surname", "family name"))
    if first and last:
        return f"[{last}]_[{first}]"
    email = guess_email_column(headers)
    if email:
        return f"[{email}]"
    if headers:
        return f"[{headers[0]}]"
    return "draft_[Row]"


def record_title(record: dict[str, str], headers: list[str]) -> str:
    first = find_header(headers, ("first name", "firstname", "given name"))
    last = find_header(headers, ("last name", "lastname", "surname", "family name"))
    if first or last:
        name = " ".join(
            part
            for part in (
                record.get(first or "", ""),
                record.get(last or "", ""),
            )
            if part
        ).strip()
        if name:
            return name
    for header in headers:
        if record.get(header):
            return record[header]
    return "Draft"


def safe_stem(value: str) -> str:
    cleaned = _INVALID_FILENAME.sub(" ", value)
    cleaned = re.sub(r"\s+", " ", cleaned).strip().rstrip(". ")
    cleaned = cleaned[:80].strip().rstrip(". ")
    if cleaned.upper() in _RESERVED:
        cleaned = f"_{cleaned}"
    if cleaned in {"", ".", ".."}:
        return ""
    return cleaned


def load_table(path: Path | str, sheet: str | None = None) -> Table:
    path = Path(path)
    if not path.is_file():
        raise FileNotFoundError(f"Could not find the file {path}.")
    suffix = path.suffix.lower()
    if suffix == ".csv":
        rows = _load_csv(path)
        sheets: list[str] = []
        chosen = None
    elif suffix in {".xlsx", ".xlsm"}:
        sheets, rows, chosen = _load_excel(path, sheet)
    elif suffix == ".xls":
        raise ValueError(
            "Composer reads .xlsx, .xlsm, and .csv files. In Excel, use Save As and choose Excel Workbook (.xlsx)."
        )
    else:
        raise ValueError("Choose a .xlsx, .xlsm, or .csv file.")
    headers, records = records_from_rows(rows)
    if not headers:
        raise ValueError("The spreadsheet does not have any columns.")
    if not records:
        raise ValueError("The spreadsheet has column names but no data rows.")
    return Table(headers=headers, records=records, sheets=sheets, sheet=chosen, source=path)


def records_from_rows(rows: list[list[str]]) -> tuple[list[str], list[dict[str, str]]]:
    pending = [list(row) for row in rows]
    while pending and not any(cell.strip() for cell in pending[0]):
        pending.pop(0)
    if not pending:
        return [], []
    raw_headers = pending[0]
    data_rows = pending[1:]
    width = max([len(raw_headers), *(len(row) for row in data_rows)] or [0])
    seen: set[str] = set()
    columns: list[tuple[str, bool]] = []
    for index in range(width):
        original = raw_headers[index].strip() if index < len(raw_headers) else ""
        synthetic = not original
        name = original or f"Column {index + 1}"
        candidate = name
        number = 2
        while normalize(candidate) in seen:
            candidate = f"{name} ({number})"
            number += 1
        seen.add(normalize(candidate))
        columns.append((candidate, synthetic))
    records: list[dict[str, str]] = []
    for row in data_rows:
        values = list(row) + [""] * (width - len(row))
        if all(not cell.strip() for cell in values[:width]):
            continue
        records.append({columns[index][0]: values[index].strip() for index in range(width)})
    kept = [
        name
        for name, synthetic in columns
        if not (synthetic and all(record.get(name, "") == "" for record in records))
    ]
    if not kept:
        return [], records
    trimmed = [{name: record.get(name, "") for name in kept} for record in records]
    return kept, trimmed


def cell_to_str(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, datetime):
        if value.hour == value.minute == value.second == value.microsecond == 0:
            return value.strftime("%Y-%m-%d")
        return value.strftime("%Y-%m-%d %H:%M")
    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, time):
        return value.strftime("%H:%M")
    if isinstance(value, float):
        if math.isnan(value) or math.isinf(value):
            return ""
        if value.is_integer() and abs(value) < 10**15:
            return str(int(value))
        return f"{value:.10f}".rstrip("0").rstrip(".")
    if isinstance(value, int):
        return str(value)
    return str(value).strip()


def _load_csv(path: Path) -> list[list[str]]:
    raw = path.read_bytes()
    if not raw.strip():
        raise ValueError(f"{path.name} is empty.")
    text = None
    for encoding in ("utf-8-sig", "utf-8", "cp1252"):
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        text = raw.decode("latin-1")
    return parse_csv(text)


def parse_csv(text: str) -> list[list[str]]:
    sample = text[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    rows = list(csv.reader(io.StringIO(text), dialect))
    if rows and max((len(row) for row in rows[:5]), default=1) == 1:
        first_cells = [row[0] for row in rows[:5] if row]
        if any(";" in cell for cell in first_cells):
            rows = list(csv.reader(io.StringIO(text), delimiter=";"))
    return rows


def _load_excel(path: Path, sheet: str | None) -> tuple[list[str], list[list[str]], str]:
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(filename=path, read_only=True, data_only=True)
    except PermissionError as exc:
        raise PermissionError(
            f"Could not open {path.name}. If it is open in Excel, close it or save a copy, then try again."
        ) from exc
    except BadZipFile as exc:
        raise ValueError(f"{path.name} is not a valid Excel workbook.") from exc
    try:
        sheets = list(workbook.sheetnames)
        if not sheets:
            raise ValueError(f"{path.name} has no worksheets.")
        if sheet is None:
            chosen = sheets[0]
        elif sheet not in sheets:
            raise ValueError(f"{path.name} has no sheet named {sheet}.")
        else:
            chosen = sheet
        worksheet = workbook[chosen]
        rows = [[cell_to_str(value) for value in row] for row in worksheet.iter_rows(values_only=True)]
        return sheets, rows, chosen
    finally:
        workbook.close()
