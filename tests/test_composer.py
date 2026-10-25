"""Checks for merging, draft files, and the rule that nothing is sent by itself."""

from __future__ import annotations

import inspect
import os
import tempfile
import unittest
import zipfile
from datetime import date
from email import policy
from email.parser import BytesParser
from pathlib import Path

from docx import Document

import composer.export as export
from composer.cities import computed_fields, nearest_city, normalize_zip
from composer.data import (
    EXAMPLE_BODY,
    EXAMPLE_SUBJECT,
    closest_header,
    default_filename_pattern,
    guess_email_column,
    load_table,
    looks_like_email,
    parse_csv,
    render,
    safe_stem,
)
from composer.export import EMAIL, WORD, ExportJob, export_batch
from composer.mailer import SmtpConfig, send_files
from composer.paths import sample_csv
from composer.settings import load_settings, save_settings, settings_file

EXPECTED = (
    "Hello Jane Adler,\n"
    "\n"
    "We noticed your phone number is 555-0101 and your email is jane.adler@example.com. Is this correct?"
)


class MergeTests(unittest.TestCase):
    def test_sample_sentence_matches_the_requested_template(self) -> None:
        table = load_table(sample_csv())
        self.assertEqual(
            table.headers,
            ["first name", "last name", "ip", "phone", "email"],
        )
        merged = render(EXAMPLE_BODY, table.records[0], table.headers, 1)
        self.assertEqual(merged.text, EXPECTED)
        self.assertEqual(merged.unknown, [])
        subject = render(EXAMPLE_SUBJECT, table.records[0], table.headers, 1)
        self.assertEqual(subject.text, "Hello Jane")

    def test_matching_ignores_case_and_extra_spaces(self) -> None:
        headers = ["First   Name"]
        merged = render("Hi [first name]", {"First   Name": "Ada"}, headers, 1)
        self.assertEqual(merged.text, "Hi Ada")
        self.assertEqual(merged.unknown, [])

    def test_unknown_token_is_left_in_place(self) -> None:
        merged = render("Hi [Nickname]", {"First Name": "Ada"}, ["First Name"], 4)
        self.assertEqual(merged.text, "Hi [Nickname]")
        self.assertEqual(merged.unknown, ["[Nickname]"])
        self.assertEqual(closest_header("Nickname", ["First Name"]), None)
        self.assertEqual(closest_header("Phone", ["Phone Number"]), "Phone Number")

    def test_nearest_city_is_a_field(self) -> None:
        self.assertEqual(normalize_zip("90210-1234"), "90210")
        self.assertEqual(normalize_zip("501"), "00501")
        self.assertEqual(nearest_city("10001"), "New York, NY")
        self.assertEqual(nearest_city("02101"), "Boston, MA")
        self.assertEqual(nearest_city("90210"), "Los Angeles, CA")
        self.assertEqual(nearest_city("77002"), "Houston, TX")
        self.assertEqual(nearest_city(""), "")
        extras = computed_fields({"zip code": "10001"}, ["zip code"], "zip code")
        merged = render("Hello from [Nearest City].", {"zip code": "10001"}, ["zip code"], 1, extras)
        self.assertEqual(merged.text, "Hello from New York, NY.")
        self.assertEqual(merged.unknown, [])

    def test_row_token_and_empty_cell(self) -> None:
        merged = render("[Row]-[Phone]", {"Phone": ""}, ["Phone"], 12)
        self.assertEqual(merged.text, "12-")

    def test_safe_filename(self) -> None:
        self.assertEqual(safe_stem("Jane:Adler"), "Jane Adler")
        self.assertEqual(safe_stem("CON"), "_CON")
        self.assertEqual(safe_stem("  ...  "), "")
        self.assertEqual(safe_stem("a/b"), "a b")

    def test_guess_email_column(self) -> None:
        self.assertEqual(guess_email_column(["first name", "e-mail"]), "e-mail")
        self.assertEqual(guess_email_column(["Notes", "Email Address"]), "Email Address")
        self.assertIsNone(guess_email_column(["phone", "ip"]))

    def test_email_shape(self) -> None:
        self.assertTrue(looks_like_email("jane.adler@example.com"))
        self.assertFalse(looks_like_email("not an email"))
        self.assertFalse(looks_like_email("a@b"))
        self.assertFalse(looks_like_email(""))

    def test_semicolon_and_quotes(self) -> None:
        rows = parse_csv('first name;notes\n"Ada, M.";"hello; there"\n')
        self.assertEqual(rows[0], ["first name", "notes"])
        self.assertEqual(rows[1][0], "Ada, M.")

    def test_excel_values_and_sheet(self) -> None:
        from openpyxl import Workbook

        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "people.xlsx"
            workbook = Workbook()
            first = workbook.active
            first.title = "People"
            first.append(["First Name", "Phone", "Email", "Start"])
            first.append(["Jane", 5550101, "jane.adler@example.com", date(2026, 10, 24)])
            other = workbook.create_sheet("Other")
            other.append(["First Name"])
            other.append(["Skipped"])
            workbook.save(path)
            table = load_table(path, sheet="People")
            self.assertEqual(table.sheet, "People")
            self.assertEqual(table.records[0]["Phone"], "5550101")
            self.assertEqual(table.records[0]["Start"], "2026-10-24")
            self.assertEqual(table.records[0]["First Name"], "Jane")


class ExportTests(unittest.TestCase):
    def setUp(self) -> None:
        self.tmp = tempfile.TemporaryDirectory()
        self.output = Path(self.tmp.name)

    def tearDown(self) -> None:
        self.tmp.cleanup()

    def test_export_module_cannot_send(self) -> None:
        source = inspect.getsource(export)
        self.assertNotIn("smtplib", source)
        self.assertNotIn("send_files", source)

    def test_sample_emails_are_drafts(self) -> None:
        table = load_table(sample_csv())
        report = export_batch(
            ExportJob(
                kind=EMAIL,
                headers=table.headers,
                records=table.records,
                source_name=table.source.name,
                output_dir=self.output,
                subject_template=EXAMPLE_SUBJECT,
                body_template=EXAMPLE_BODY,
                filename_template=default_filename_pattern(table.headers),
                recipient_column="email",
                from_name="Composer",
                from_email="reviews@example.com",
            )
        )
        self.assertEqual(len(report.written), 6)
        self.assertEqual(report.skipped, [])
        notes = (report.folder / "Review notes.txt").read_text(encoding="utf-8")
        self.assertIn("Composer did not send any email.", notes)
        message = BytesParser(policy=policy.default).parsebytes((report.folder / "Adler_Jane.eml").read_bytes())
        body = message.get_content().replace("\r\n", "\n").strip()
        self.assertEqual(body, EXPECTED)
        self.assertEqual(str(message["Subject"]), "Hello Jane")
        self.assertEqual(str(message["To"]), "jane.adler@example.com")
        self.assertEqual(str(message["X-Unsent"]), "1")

    def test_subject_newlines_cannot_add_headers(self) -> None:
        report = export_batch(
            ExportJob(
                kind=EMAIL,
                headers=["email"],
                records=[{"email": "ada@example.com"}],
                source_name="one.csv",
                output_dir=self.output,
                subject_template="Hello\nBcc: evil@example.com",
                body_template="Hi",
                filename_template="note",
                recipient_column="email",
            )
        )
        message = BytesParser(policy=policy.default).parsebytes(report.written[0].path.read_bytes())
        self.assertNotIn("Bcc", message)
        self.assertEqual(str(message["Subject"]), "Hello Bcc: evil@example.com")

    def test_invalid_email_is_skipped(self) -> None:
        report = export_batch(
            ExportJob(
                kind=EMAIL,
                headers=["email", "first name"],
                records=[
                    {"email": "ada@example.com", "first name": "Ada"},
                    {"email": "", "first name": "No"},
                    {"email": "not-an-email", "first name": "Bad"},
                ],
                source_name="mixed.csv",
                output_dir=self.output,
                subject_template="Hi",
                body_template="Hello [first name]",
                filename_template="[first name]",
                recipient_column="email",
            )
        )
        self.assertEqual(len(report.written), 1)
        self.assertEqual(len(report.skipped), 2)
        self.assertIn("no email address", report.skipped[0][1])

    def test_duplicate_names_get_a_suffix(self) -> None:
        report = export_batch(
            ExportJob(
                kind=WORD,
                headers=["first name"],
                records=[{"first name": "Ada"}, {"first name": "Ada"}],
                source_name="dupes.csv",
                output_dir=self.output,
                subject_template="",
                body_template="Hello [first name]",
                filename_template="[first name]",
                recipient_column="",
            )
        )
        names = sorted(draft.filename for draft in report.written)
        self.assertEqual(names, ["Ada.docx", "Ada_2.docx"])
        text = "\n".join(paragraph.text for paragraph in Document(report.written[0].path).paragraphs)
        self.assertIn("Hello Ada", text)

    def test_create_zip_packs_the_generated_files(self) -> None:
        report = export_batch(
            ExportJob(
                kind=EMAIL,
                headers=["email"],
                records=[{"email": "ada@example.com"}],
                source_name="one.csv",
                output_dir=self.output,
                subject_template="Hi",
                body_template="Hello",
                filename_template="note",
                recipient_column="email",
                create_zip=True,
            )
        )
        self.assertIsNotNone(report.archive)
        self.assertTrue(report.archive.is_file())
        self.assertEqual(report.archive.parent, report.folder.parent)
        with zipfile.ZipFile(report.archive) as archive:
            names = archive.namelist()
        self.assertTrue(any(name.endswith("note.eml") for name in names))
        self.assertTrue(any(name.endswith("Review notes.txt") for name in names))

    def test_cancel_writes_nothing(self) -> None:
        report = export_batch(
            ExportJob(
                kind=EMAIL,
                headers=["email"],
                records=[{"email": "ada@example.com"}],
                source_name="one.csv",
                output_dir=self.output,
                subject_template="Hi",
                body_template="Hello",
                filename_template="note",
                recipient_column="email",
            ),
            cancel=lambda: True,
        )
        self.assertTrue(report.cancelled)
        self.assertEqual(report.written, [])

    def test_send_refuses_without_confirmation(self) -> None:
        with self.assertRaises(RuntimeError):
            send_files(
                [],
                SmtpConfig("localhost", 25, "none", "", "", "", "ada@example.com"),
                confirmed=False,
            )


class SettingsTests(unittest.TestCase):
    def test_password_is_not_stored(self) -> None:
        with tempfile.TemporaryDirectory() as folder:
            os.environ["COMPOSER_SETTINGS_DIR"] = folder
            save_settings(
                {
                    "subject": "Hi",
                    "password": "secret",
                    "smtp_password": "secret2",
                    "from_email": "ada@example.com",
                }
            )
            raw = settings_file().read_text(encoding="utf-8")
            self.assertNotIn("secret", raw)
            self.assertNotIn("password", raw)
            loaded = load_settings()
            self.assertEqual(loaded["subject"], "Hi")
            self.assertEqual(loaded["from_email"], "ada@example.com")


@unittest.skipUnless(__import__("importlib").util.find_spec("PySide6"), "PySide6 is not installed")
class GuiTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        os.environ["QT_QPA_PLATFORM"] = "offscreen"
        from PySide6.QtWidgets import QApplication

        cls.app = QApplication.instance() or QApplication([])

    def setUp(self) -> None:
        self.tmp = tempfile.TemporaryDirectory()
        os.environ["COMPOSER_SETTINGS_DIR"] = self.tmp.name

    def tearDown(self) -> None:
        self.tmp.cleanup()

    def test_preview_fills_the_sample_message(self) -> None:
        from composer.window import MainWindow

        window = MainWindow()
        window.load_file(sample_csv())
        self.assertIn("Hello Jane Adler,", window.preview_plain())
        self.assertIn("555-0101", window.preview_plain())
        self.assertIn("jane.adler@example.com", window.preview_plain())
        self.assertIn("Hello Jane", window.preview_subject())
        window.set_kind("word")
        self.assertFalse(window.subject_row.isVisible())
        self.assertNotIn("sent", window.status_label.text().casefold())
        window.close()
        self.app.processEvents()


if __name__ == "__main__":
    unittest.main()
